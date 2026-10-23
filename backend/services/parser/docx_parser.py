"""Word 简历解析（v2 · 阶段 1 实现）

- 使用 python-docx 提取段落 + 表格文本，输出纯文本 + 分段；
- 表格按行拼接为 `单元格 | 单元格` 文本（简历里的时间线/技能表常见）；
- 解析失败抛 `ParserError`（API 层映射 code=1002）。
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from .base import (
    ParsedResume,
    ParserError,
    build_segments,
    looks_garbled,
    read_source,
)


def parse_docx(
    source: str | Path | bytes,
    *,
    filename: str | None = None,
) -> ParsedResume:
    """解析 DOCX 简历。

    Args:
        source: 文件路径或原始字节（便于 API 层直接传 UploadFile 内容）。
        filename: 覆盖返回结果中的文件名（source 为 bytes 时使用）。
    """
    data, name = read_source(source, filename=filename)
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - venv 已安装 python-docx
        raise ParserError("缺少 python-docx 依赖，请执行 pip install -r requirements.txt") from exc

    try:
        document = Document(BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - 损坏 / 加密 / 非法格式统一报解析失败
        raise ParserError(f"DOCX 解析失败（文件可能损坏或加密）：{exc}") from exc

    paragraphs: list[str] = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text and paragraph.text.strip()
    ]

    table_rows: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                table_rows.append(" | ".join(cells))

    body_text = "\n".join(paragraphs)
    table_text = "\n".join(table_rows)
    text = "\n".join(part for part in (body_text, table_text) if part).strip()
    if not text:
        raise ParserError("DOCX 中未提取到文本（可能为纯图片文档，请改用粘贴文本）")

    warnings: list[str] = []
    degraded = False
    if looks_garbled(text):
        degraded = True
        warnings.append("提取文本疑似乱码，建议核对或改用粘贴文本输入")
    if not paragraphs and table_rows:
        warnings.append("文档主体为表格，已提取为纯文本")

    return ParsedResume(
        filename=name,
        kind="docx",
        text=text,
        segments=build_segments(paragraphs),
        page_count=None,
        degraded=degraded,
        warnings=warnings,
    )


def parse_docx_stream(fileobj: Any) -> ParsedResume:
    """从文件对象解析（FastAPI UploadFile.file 可直接传入）。"""
    return parse_docx(fileobj.read(), filename=getattr(fileobj, "name", "resume.docx"))


__all__ = ["parse_docx", "parse_docx_stream"]
