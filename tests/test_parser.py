"""简历解析测试（v2 · 阶段 1 实现）

覆盖：
- DOCX：段落 + 表格提取、小节分段、空文档报错
- PDF：pdfplumber 主链路（注入提取器验证表格拼接）、扫描件降级、损坏文件报错
- 统一入口 parse_resume：格式分派 / TXT / 文件不存在 / 空内容 / 不支持格式

不依赖网络与真实文件（DOCX / PDF 均在内存中生成）。
"""
from __future__ import annotations

from io import BytesIO

import pytest

from backend.services.parser import (
    ParsedResume,
    ParserError,
    UnsupportedResumeFormatError,
    parse_resume,
)
from backend.services.parser.pdf_parser import parse_pdf


def make_docx_bytes() -> bytes:
    """构造含段落 + 表格的 DOCX（python-docx 内存生成）。"""
    from docx import Document

    document = Document()
    document.add_paragraph("教育经历")
    document.add_paragraph("某大学 计算机科学与技术 2022–2026")
    document.add_paragraph("项目经历")
    document.add_paragraph("AI 面试助手：负责后端 FastAPI 接口开发")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "技能"
    table.cell(0, 1).text = "Python"
    table.cell(1, 0).text = "证书"
    table.cell(1, 1).text = "CET-6"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_blank_pdf_bytes() -> bytes:
    """构造一页空白 PDF（pypdf 内存生成，用于扫描件降级测试）。"""
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ #
# DOCX
# ------------------------------------------------------------------ #
def test_parse_docx_extracts_paragraphs_tables_and_segments():
    doc = parse_resume("resume.docx", make_docx_bytes())
    assert isinstance(doc, ParsedResume)
    assert doc.kind == "docx"
    assert doc.degraded is False
    assert "某大学 计算机科学与技术" in doc.text
    assert "Python" in doc.text
    assert "|" in doc.text  # 表格文本拼接
    headings = [segment.heading for segment in doc.segments]
    assert "教育经历" in headings
    assert "项目经历" in headings
    assert doc.char_count == len(doc.text)
    assert doc.page_count is None

    payload = doc.to_dict()
    assert payload["filename"] == "resume.docx"
    assert payload["segments"][0]["heading"] == "教育经历"
    assert "degraded" in payload and "warnings" in payload


def test_parse_docx_empty_document_raises():
    from docx import Document

    document = Document()  # 无任何文本
    buffer = BytesIO()
    document.save(buffer)
    with pytest.raises(ParserError, match="未提取到文本"):
        parse_resume("empty.docx", buffer.getvalue())


def test_parse_docx_corrupted_bytes_raises():
    with pytest.raises(ParserError, match="DOCX 解析失败"):
        parse_resume("broken.docx", b"not a real docx")


# ------------------------------------------------------------------ #
# PDF
# ------------------------------------------------------------------ #
def test_parse_pdf_blank_page_marks_degraded():
    doc = parse_resume("scan.pdf", make_blank_pdf_bytes())
    assert doc.kind == "pdf"
    assert doc.page_count == 1
    assert doc.degraded is True
    assert doc.text == ""
    assert any("扫描件" in warning or "未提取" in warning for warning in doc.warnings)


def test_parse_pdf_with_injected_extractor_combines_text_and_tables():
    def fake_extractor(_data: bytes):
        return (
            [
                "教育经历\n某大学 计算机科学与技术 2022–2026\n"
                "项目经历\nAI 面试助手：负责后端 FastAPI 接口开发与测试"
            ],
            ["技能 | Python", "证书 | CET-6"],
        )

    doc = parse_pdf(b"%PDF-fake", extractor=fake_extractor)
    assert doc.kind == "pdf"
    assert doc.page_count == 1
    assert doc.degraded is False
    assert "某大学" in doc.text
    assert "技能 | Python" in doc.text
    assert doc.segments[0].heading == "教育经历"


def test_parse_pdf_no_pages_raises():
    with pytest.raises(ParserError, match="没有可解析的页面"):
        parse_pdf(b"%PDF-fake", extractor=lambda _data: ([], []))


def test_parse_pdf_scanned_partial_pages_marks_degraded():
    doc = parse_pdf(b"%PDF-fake", extractor=lambda _data: (["", ""], []))
    assert doc.degraded is True
    assert doc.page_count == 2
    assert any("扫描件" in warning for warning in doc.warnings)


def test_parse_pdf_garbled_text_marks_degraded():
    garbled = "\x00\x01\x02\x03" * 100
    doc = parse_pdf(b"%PDF-fake", extractor=lambda _data: ([garbled], []))
    assert doc.degraded is True
    assert any("乱码" in warning for warning in doc.warnings)


def test_parse_pdf_corrupted_bytes_raises():
    with pytest.raises(ParserError):
        parse_pdf(b"definitely not a pdf")


# ------------------------------------------------------------------ #
# 统一入口 parse_resume
# ------------------------------------------------------------------ #
def test_parse_resume_txt():
    doc = parse_resume("resume.txt", "教育经历\n某大学 计算机".encode("utf-8"))
    assert doc.kind == "text"
    assert "某大学" in doc.text
    assert doc.segments[0].heading == "教育经历"


def test_parse_resume_unsupported_suffix():
    with pytest.raises(UnsupportedResumeFormatError, match="不支持的文件格式"):
        parse_resume("resume.xlsx", b"whatever")


def test_parse_resume_missing_file():
    with pytest.raises(ParserError, match="文件不存在"):
        parse_resume("no-such-resume.pdf")


def test_parse_resume_empty_bytes():
    with pytest.raises(ParserError, match="文件内容为空"):
        parse_resume("resume.docx", b"")
